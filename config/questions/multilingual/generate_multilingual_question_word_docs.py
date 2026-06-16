#!/usr/bin/env python3
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


BASE_DIR = Path(__file__).resolve().parent
DOCS_DIR = BASE_DIR / "multilingual_questions_docs"
JSON_DIR = BASE_DIR / "multilingual_questions_json"
COMPLETE_JSON = BASE_DIR / "multilingual_questions_complete.json"

QUESTION_ORDER = [
    "A008",
    "A165",
    "E018",
    "E025",
    "F063",
    "F118",
    "F120",
    "G006",
    "Y002",
    "Y003",
]

ITEM_LABELS = {
    "A008": "Feeling of happiness.",
    "A165": "Most people can be trusted.",
    "E018": "Greater respect for authority.",
    "E025": "Political action: signing a petition.",
    "F063": "How important is God in your life.",
    "F118": "Justifiable: homosexuality.",
    "F120": "Justifiable: abortion.",
    "G006": "How proud of nationality.",
    "Y002": "Post-Materialist index (4-item).",
    "Y003": "Autonomy Index.",
}

LANGUAGE_NAMES = {
    "ar": "Arabic",
    "ru": "Russian",
    "zh-cn": "Chinese Simplified",
    "zh-tw": "Chinese Traditional (Taiwan)",
    "zh-hk": "Cantonese (Hong Kong)",
}

LANGUAGE_FONTS = {
    "ar": ("Arial", "Arial"),
    "ru": ("Times New Roman", "Times New Roman"),
    "zh-cn": ("Times New Roman", "SimSun"),
    "zh-tw": ("Times New Roman", "PMingLiU"),
    "zh-hk": ("Times New Roman", "PMingLiU"),
}


def normalize(text: str) -> str:
    return " ".join(text.replace("\r", "\n").replace("\n", " ").split())


def set_run_fonts(run, latin_font: str, east_asia_font: str) -> None:
    run.font.name = latin_font
    run.font.size = Pt(12)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), latin_font)
    rfonts.set(qn("w:hAnsi"), latin_font)
    rfonts.set(qn("w:cs"), latin_font)
    rfonts.set(qn("w:eastAsia"), east_asia_font)


def set_run_rtl(run) -> None:
    rpr = run._element.get_or_add_rPr()
    rtl = rpr.find(qn("w:rtl"))
    if rtl is None:
        rtl = OxmlElement("w:rtl")
        rpr.append(rtl)
    rtl.set(qn("w:val"), "1")


def set_paragraph_bidi(paragraph) -> None:
    ppr = paragraph._element.get_or_add_pPr()
    bidi = ppr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        ppr.append(bidi)
    bidi.set(qn("w:val"), "1")


def load_standard_language(language_code: str) -> dict:
    path = JSON_DIR / f"questions_{language_code}.json"
    return json.loads(path.read_text(encoding="utf-8"))


def load_complete_language(language_code: str) -> dict:
    data = json.loads(COMPLETE_JSON.read_text(encoding="utf-8"))
    lang = data["languages"][language_code]
    out = {}
    for code, entry in lang["questions"].items():
        out[code] = {
            "question": entry["question"],
            "scale": entry["scale"],
        }
    return out


def load_language(language_code: str) -> dict:
    if language_code in {"zh-tw", "zh-hk"}:
        return load_complete_language(language_code)
    return load_standard_language(language_code)


def style_document(document: Document, language_code: str) -> None:
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    latin_font, east_asia_font = LANGUAGE_FONTS[language_code]
    normal = document.styles["Normal"]
    normal.font.name = latin_font
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:ascii"), latin_font)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), latin_font)
    normal._element.rPr.rFonts.set(qn("w:cs"), latin_font)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia_font)


def add_title(document: Document, language_code: str) -> None:
    latin_font, east_asia_font = LANGUAGE_FONTS[language_code]
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if language_code == "ar" else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(LANGUAGE_NAMES[language_code])
    run.bold = True
    set_run_fonts(run, latin_font, east_asia_font)
    if language_code == "ar":
        set_paragraph_bidi(p)
        set_run_rtl(run)


def add_question(document: Document, language_code: str, code: str, question: str) -> None:
    latin_font, east_asia_font = LANGUAGE_FONTS[language_code]
    if language_code == "ar":
        p_label = document.add_paragraph()
        fmt_label = p_label.paragraph_format
        fmt_label.space_after = Pt(4)
        fmt_label.line_spacing = 1.15
        p_label.alignment = WD_ALIGN_PARAGRAPH.LEFT

        run_code = p_label.add_run(f"{code}  ")
        run_code.bold = True
        set_run_fonts(run_code, latin_font, east_asia_font)

        run_label = p_label.add_run(f"{ITEM_LABELS[code]}")
        run_label.italic = True
        set_run_fonts(run_label, latin_font, east_asia_font)

        p_question = document.add_paragraph()
        fmt_q = p_question.paragraph_format
        fmt_q.right_indent = Inches(0.35)
        fmt_q.left_indent = Inches(0.15)
        fmt_q.space_after = Pt(12)
        fmt_q.line_spacing = 1.35
        p_question.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_paragraph_bidi(p_question)

        run_question = p_question.add_run(normalize(question))
        set_run_fonts(run_question, latin_font, east_asia_font)
        set_run_rtl(run_question)
        return

    p = document.add_paragraph()
    fmt = p.paragraph_format
    fmt.left_indent = Inches(0.45)
    fmt.first_line_indent = Inches(-0.45)
    fmt.space_after = Pt(12)
    fmt.line_spacing = 1.35

    run_code = p.add_run(f"{code}  ")
    run_code.bold = True
    set_run_fonts(run_code, latin_font, east_asia_font)

    run_label = p.add_run(f"{ITEM_LABELS[code]} ")
    run_label.italic = True
    set_run_fonts(run_label, latin_font, east_asia_font)

    run_question = p.add_run(f'"{normalize(question)}"')
    set_run_fonts(run_question, latin_font, east_asia_font)


def build_docx(language_code: str, questions: dict) -> Path:
    document = Document()
    style_document(document, language_code)
    add_title(document, language_code)
    for code in QUESTION_ORDER:
        entry = questions.get(code)
        question_text = entry["question"] if entry else "[Missing in source]"
        add_question(document, language_code, code, question_text)
    output = DOCS_DIR / f"questions_{language_code}.docx"
    document.save(output)
    return output


def build_combined_docx(language_codes: list[str]) -> Path:
    document = Document()
    first = True
    for language_code in language_codes:
        questions = load_language(language_code)
        if first:
            style_document(document, language_code)
            first = False
        else:
            document.add_page_break()
        add_title(document, language_code)
        for code in QUESTION_ORDER:
            entry = questions.get(code)
            question_text = entry["question"] if entry else "[Missing in source]"
            add_question(document, language_code, code, question_text)
    output = DOCS_DIR / "questions_selected_languages_combined.docx"
    document.save(output)
    return output


def main() -> None:
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    language_codes = ["ar", "ru", "zh-cn", "zh-tw", "zh-hk"]
    for language_code in language_codes:
        path = build_docx(language_code, load_language(language_code))
        print(f"Wrote {path}")
    combined = build_combined_docx(language_codes)
    print(f"Wrote {combined}")


if __name__ == "__main__":
    main()
